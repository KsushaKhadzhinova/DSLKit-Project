from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

screens = Path(__file__).parent / "screens"
GREEN, RED = "C6EFCE", "FFC7CE"


def new_doc(title, intro):
    d = Document()
    s = d.sections[0]
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin = s.bottom_margin = Cm(1.8)
    for name, size in (("Normal", 14), ("Heading 1", 16), ("Heading 2", 15), ("Heading 3", 14)):
        st = d.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            st.element.rPr.rFonts.attrib.pop(qn(attr), None)
        if name != "Normal":
            st.font.bold = True
    d.add_paragraph(intro)
    return d


def shade(cell, color):
    tc = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), color)
    tc.append(el)


def table(d, head, rows, widths=None, status_col=None):
    t = d.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].add_run(h).bold = True
        shade(c, "D9E2F3")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
        if status_col is not None:
            s = str(r[status_col])
            shade(cells[status_col], GREEN if s.startswith("Passed") or "☑" in s else RED)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(12)
    d.add_paragraph()
    return t


def shot(d, name, caption, width=16):
    d.add_picture(str(screens / name), width=Cm(width))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True


def rich(par, text, size=12):
    import re
    for i, line in enumerate(text.split("\n")):
        if i:
            par.add_run().add_break()
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", line):
            if not part:
                continue
            bold, ital = part.startswith("**"), part.startswith("*") and not part.startswith("**")
            run = par.add_run(part.strip("*") if (bold or ital) else part)
            run.bold, run.italic = bold or None, ital or None
            run.font.size = Pt(size)


def use_case(d, path):
    import json
    base = Path(path).parent
    rows = json.loads(Path(path).read_text(encoding="utf8"))
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        c = t.add_row().cells
        c[0].text = ""
        c[0].paragraphs[0].add_run(k).bold = True
        c[0].paragraphs[0].runs[0].font.size = Pt(12)
        rich(c[1].paragraphs[0], v)
        c[0].width, c[1].width = Cm(4.2), Cm(13)
        shade(c[0], "F2F2F2")
    d.add_paragraph()
    d.add_heading("Данные формы: типы и допустимые значения", 2)
    fields = json.loads((base / "fields.json").read_text(encoding="utf8"))
    table(d, ["Поле", "Тип данных", "Обяза-тельно", "По умолчанию", "Допустимые значения", "Недопустимые значения"], fields, [2.6, 2.6, 1.8, 2, 4.2, 4.2])
    d.add_heading("Диаграммы процесса", 2)
    for name, cap, w in (("flow", "Рис. A. Блок-схема сценария UC-01: основной поток и альтернативные потоки", 10),
                         ("states", "Рис. B. Состояния поля формы при проверке", 16),
                         ("components", "Рис. C. Компоненты и обмен данными при отправке заявки", 16)):
        d.add_picture(str(base / "diagrams" / f"{name}.png"), width=Cm(w))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = d.add_paragraph(cap)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True
